from __future__ import annotations

import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from converter import (
    ANALYSIS_FIELDNAMES,
    convert_file,
    iter_analysis_rows,
    output_paths,
    paths_for_directory,
    validate_input_file,
)
from decoder import FIELDNAMES, TIMESTAMP_MODULUS, TimestampUnwrapper, decode_record


def make_record(*, timestamp: int, control: int, packet_id: int, payload: int) -> bytes:
    header = (
        (timestamp & (TIMESTAMP_MODULUS - 1))
        | ((control & 0x07) << 21)
        | ((packet_id & 0xFF) << 24)
    )
    return header.to_bytes(4, "little") + (payload & 0xFFFFFFFF).to_bytes(4, "little")


class InputValidationTests(unittest.TestCase):
    def test_empty_file_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "empty.bin"
            path.write_bytes(b"")
            with self.assertRaisesRegex(ValueError, "vazio"):
                validate_input_file(path)

    def test_non_multiple_of_eight_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "invalid.bin"
            path.write_bytes(b"123456789")
            with self.assertRaisesRegex(ValueError, "multiplo de 8"):
                validate_input_file(path)


class OutputPathTests(unittest.TestCase):
    def test_default_output_paths(self) -> None:
        input_path = Path("C:/dados/data001.bin")
        csv_path, xlsx_path = output_paths(input_path, None, "both")
        self.assertEqual(csv_path, Path("C:/dados/data001_decoded.csv"))
        self.assertEqual(xlsx_path, Path("C:/dados/data001_decoded.xlsx"))

    def test_gui_directory_output_paths(self) -> None:
        input_path = Path("C:/dados/data001.bin")
        output_dir = Path("C:/resultado")
        csv_path, xlsx_path = paths_for_directory(input_path, output_dir, "xlsx")
        self.assertIsNone(csv_path)
        self.assertEqual(xlsx_path, Path("C:/resultado/data001_decoded.xlsx"))

    def test_debug_output_paths_have_a_distinct_name(self) -> None:
        input_path = Path("C:/dados/data001.bin")
        output_dir = Path("C:/resultado")

        default_csv, default_xlsx = output_paths(
            input_path, None, "both", extended_debug=True
        )
        gui_csv, gui_xlsx = paths_for_directory(
            input_path, output_dir, "both", extended_debug=True
        )

        self.assertEqual(default_csv, Path("C:/dados/data001_decoded_debug.csv"))
        self.assertEqual(default_xlsx, Path("C:/dados/data001_decoded_debug.xlsx"))
        self.assertEqual(gui_csv, Path("C:/resultado/data001_decoded_debug.csv"))
        self.assertEqual(gui_xlsx, Path("C:/resultado/data001_decoded_debug.xlsx"))


class CsvConversionTests(unittest.TestCase):
    def test_csv_conversion_and_progress(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            input_path = temp / "data001.bin"
            # SESSION_MARKER conforme a logica existente: header/payload zerados.
            input_path.write_bytes(bytes(8) + bytes(8))
            csv_path = temp / "data001_decoded.csv"
            progress: list[tuple[int, int]] = []

            result = convert_file(
                input_path,
                csv_path=csv_path,
                xlsx_path=None,
                delimiter=";",
                progress_callback=lambda current, total: progress.append((current, total)),
                progress_interval_records=1,
            )

            self.assertEqual(result.records, 2)
            self.assertTrue(csv_path.is_file())
            self.assertEqual(progress[0], (0, 2))
            self.assertEqual(progress[-1], (2, 2))

            with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle, delimiter=";")
                rows = list(reader)
            self.assertEqual(reader.fieldnames, ANALYSIS_FIELDNAMES)
            self.assertEqual(len(rows), 2)

    def test_extended_debug_preserves_original_columns_and_one_row_per_record(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            input_path = temp / "data001.bin"
            input_path.write_bytes(bytes(16))
            csv_path = temp / "data001_debug.csv"

            result = convert_file(
                input_path,
                csv_path=csv_path,
                xlsx_path=None,
                extended_debug=True,
            )

            with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle, delimiter=";")
                rows = list(reader)
            self.assertEqual(reader.fieldnames, FIELDNAMES)
            self.assertEqual(len(rows), 2)
            self.assertEqual(result.counts["output_rows"], 2)


class AnalysisRowsTests(unittest.TestCase):
    def test_velocity_rpm_and_fuel_become_three_equivalent_signals(self) -> None:
        velocity = 104
        rpm = 2500
        fuel = 511
        payload = (velocity << 22) | (rpm << 10) | fuel
        decoded = decode_record(
            make_record(timestamp=1234, control=2, packet_id=0x01, payload=payload),
            source_file="DATA006.BIN",
            record_index=1,
            timestamp=TimestampUnwrapper(),
        )

        rows = list(iter_analysis_rows(decoded))

        self.assertEqual([row["signal"] for row in rows], ["velocity", "rpm", "fuel"])
        self.assertEqual([row["value"] for row in rows], [2.0, rpm, fuel])
        self.assertEqual([row["raw_value"] for row in rows], [velocity, rpm, fuel])
        self.assertEqual([row["unit"] for row in rows], ["m/s", "rpm", "raw_10bit"])
        self.assertTrue(all(row["source_file"] == "DATA006.BIN" for row in rows))
        self.assertTrue(all(row["byte_offset"] == 8 for row in rows))
        self.assertTrue(all(row["timestamp_raw_ms"] == 1234 for row in rows))
        self.assertTrue(all(row["time_s"] == 1.234 for row in rows))
        self.assertTrue(all(row["flags"] == "OK" for row in rows))

    def test_imu_uses_existing_logical_to_physical_axis_mapping(self) -> None:
        accel_payload = 1000
        gyro = 0xFFFF
        decoded = decode_record(
            make_record(
                timestamp=26,
                control=0,
                packet_id=0x14,
                payload=(accel_payload << 16) | gyro,
            ),
            source_file="DATA006.BIN",
            record_index=4,
            timestamp=TimestampUnwrapper(),
        )

        rows = list(iter_analysis_rows(decoded))

        self.assertEqual(decoded["record_type"], "IMU_AY_LOGICAL_GY")
        self.assertEqual([row["signal"] for row in rows], ["AY", "GY"])
        self.assertEqual([row["source_axis"] for row in rows], ["Z", "Y"])
        self.assertEqual([row["value"] for row in rows], [1001, -1])
        self.assertEqual([row["raw_value"] for row in rows], [1000, -1])
        self.assertTrue(all(row["unit"] == "LSB" for row in rows))
        self.assertTrue(all(row["flags"] == "PACKING_CORRECTED" for row in rows))

    def test_special_conditions_are_combined_in_flags(self) -> None:
        decoded = decode_record(
            make_record(timestamp=3, control=0, packet_id=0x99, payload=1),
            source_file="data001.bin",
            record_index=0,
            timestamp=TimestampUnwrapper(),
        )

        row = next(iter_analysis_rows(decoded, timestamp_wrap=True))

        self.assertEqual(row["flags"], "TIMESTAMP_WRAP|INVALID_PACKET")
        self.assertIsNone(row["signal"])
        self.assertIsNone(row["value"])

    def test_padding_is_kept_as_a_traceable_row(self) -> None:
        decoded = decode_record(
            bytes(8),
            source_file="DATA006.BIN",
            record_index=255,
            timestamp=TimestampUnwrapper(),
        )

        row = next(iter_analysis_rows(decoded))

        self.assertEqual(row["record_type"], "PADDING_CANDIDATE")
        self.assertEqual(row["byte_offset"], 2040)
        self.assertEqual(row["timestamp_raw_ms"], 0)
        self.assertIsNone(row["time_s"])
        self.assertEqual(row["flags"], "PADDING")

    def test_timestamp_wrap_flag_marks_only_the_wrap_packet(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            input_path = temp / "data001.bin"
            input_path.write_bytes(
                make_record(
                    timestamp=TIMESTAMP_MODULUS - 5,
                    control=0,
                    packet_id=0x01,
                    payload=0,
                )
                + make_record(timestamp=3, control=0, packet_id=0x01, payload=0)
            )
            csv_path = temp / "analysis.csv"

            convert_file(input_path, csv_path=csv_path, xlsx_path=None)

            with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle, delimiter=";"))
            before_wrap = [row for row in rows if row["byte_offset"] == "0"]
            after_wrap = [row for row in rows if row["byte_offset"] == "8"]
            self.assertTrue(all(row["flags"] == "OK" for row in before_wrap))
            self.assertTrue(
                all(row["flags"] == "TIMESTAMP_WRAP" for row in after_wrap)
            )
            self.assertEqual(float(after_wrap[0]["time_s"]), (TIMESTAMP_MODULUS + 3) / 1000)


class XlsxConversionTests(unittest.TestCase):
    def test_default_xlsx_is_reduced_and_debug_xlsx_is_unchanged(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            input_path = temp / "DATA006.BIN"
            velocity = 52
            rpm = 848
            fuel = 0
            payload = (velocity << 22) | (rpm << 10) | fuel
            input_path.write_bytes(
                make_record(timestamp=8, control=0, packet_id=0x01, payload=payload)
            )
            analysis_path = temp / "analysis.xlsx"
            debug_path = temp / "debug.xlsx"

            convert_file(input_path, csv_path=None, xlsx_path=analysis_path)
            convert_file(
                input_path,
                csv_path=None,
                xlsx_path=debug_path,
                extended_debug=True,
            )

            analysis_workbook = load_workbook(analysis_path, read_only=True)
            debug_workbook = load_workbook(debug_path, read_only=True)
            try:
                analysis_rows = list(
                    analysis_workbook.active.iter_rows(values_only=True)
                )
                debug_rows = list(debug_workbook.active.iter_rows(values_only=True))
            finally:
                analysis_workbook.close()
                debug_workbook.close()
            self.assertEqual(list(analysis_rows[0]), ANALYSIS_FIELDNAMES)
            self.assertEqual(list(debug_rows[0]), FIELDNAMES)
            self.assertEqual(len(analysis_rows), 4)
            self.assertEqual(len(debug_rows), 2)

            analysis_by_signal = {
                row[5]: row for row in analysis_rows[1:]
            }
            debug_header = {name: index for index, name in enumerate(debug_rows[0])}
            debug_values = debug_rows[1]
            self.assertEqual(
                analysis_by_signal["velocity"][6],
                debug_values[debug_header["velocity_m_s_estimated"]],
            )
            self.assertEqual(
                analysis_by_signal["rpm"][6],
                debug_values[debug_header["rpm_raw_12bit"]],
            )
            self.assertEqual(
                analysis_by_signal["fuel"][6],
                debug_values[debug_header["fuel_raw_10bit"]],
            )


if __name__ == "__main__":
    unittest.main()
